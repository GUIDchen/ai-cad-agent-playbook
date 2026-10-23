#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown -> DOCX converter used by this repository.

A dependency-light converter for the report in docs/: headings, pipe tables
(auto column widths, repeating header row, borders), fenced code blocks with
shading, block quotes, ordered/unordered lists, an auto generated cover page,
a table of contents and A4 page numbering in the footer.

Usage:
    pip install python-docx
    python tools/md2docx.py docs/report.md docs/report.docx

Notes:
    * The first "# title" and "## subtitle" become the cover page; following
      "> ..." lines become the cover metadata block.
    * Both triple backtick and triple tilde fences are accepted.
    * Inline markup: **bold** and `code`.
"""
import argparse
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BT = chr(96)
FENCE = BT * 3
ALT_FENCE = "~~~"

# Order of `w:shd` / `w:pBdr` successors inside `w:pPr` (OOXML sequence order).
PPR_AFTER_SHD = (
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
    "w:pPrChange",
)


def set_east_asian(obj, font):
    """Set the east-asian font on a style or run."""
    rpr = obj._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.insert_element_before(shd, *PPR_AFTER_SHD)


def add_left_border(paragraph, color="2E5496", size="12"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)
    borders.append(left)
    ppr.insert_element_before(borders, "w:shd", *PPR_AFTER_SHD)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.insert_element_before(
        shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
        "w:vAlign", "w:hideMark", "w:headers", "w:cellIns", "w:cellDel",
        "w:cellMerge", "w:tcPrChange",
    )


INLINE = re.compile(r"(\*\*.+?\*\*|" + BT + r"[^" + BT + r"]+" + BT + r")")


def add_runs(paragraph, text, size=None, color=None, code_font="Consolas",
             east_asian="宋体"):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith(BT) and part.endswith(BT) and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = code_font
            run.font.size = Pt(9.5)
            set_east_asian(run, east_asian)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def is_table_separator(line):
    line = line.strip()
    return line.startswith("|") and re.fullmatch(r"\|[\s:\-\|]+\|?", line) is not None


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def text_width(text):
    return sum(2 if ord(ch) > 127 else 1 for ch in text)


def parse_markdown(lines, start=2):
    """Return a list of (kind, payload) blocks, skipping the cover block."""
    body = []
    i = start
    while i < len(lines):
        line = lines[i].strip()
        if line in ("---", "***") or not line:
            i += 1
            continue
        match = re.match(r"^(#{1,4})\s+(.*)$", line)
        if match:
            body.append(("h", len(match.group(1)), match.group(2).strip()))
            i += 1
            continue
        if line.startswith(FENCE) or line.startswith(ALT_FENCE):
            fence = line[:3]
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith(fence):
                code.append(lines[i])
                i += 1
            i += 1
            body.append(("code", code))
            continue
        if line.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            rows = [split_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            body.append(("table", rows))
            continue
        if line.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            body.append(("quote", quote))
            continue
        if re.match(r"^[-*]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            body.append(("ul", items))
            continue
        if re.match(r"^\d+[.)]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\d+[.)]\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+[.)]\s+", "", lines[i].strip()))
                i += 1
            body.append(("ol", items))
            continue
        paragraph = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith("#") or nxt.startswith("|")
                    or nxt.startswith(">") or nxt.startswith(FENCE)
                    or nxt.startswith(ALT_FENCE) or nxt in ("---", "***")
                    or re.match(r"^[-*]\s+", nxt)
                    or re.match(r"^\d+[.)]\s+", nxt)):
                break
            paragraph.append(nxt)
            i += 1
        body.append(("p", " ".join(paragraph)))
    return body


def build(input_path, output_path, body_font="宋体", head_font="黑体",
          code_font="Consolas"):
    lines = open(input_path, encoding="utf-8").read().splitlines()
    if len(lines) < 2:
        sys.exit("input markdown needs a '# title' and a '## subtitle' line")

    title = lines[0][2:].strip()
    subtitle = lines[1][3:].strip()
    meta = []
    start = 2
    for index in range(2, len(lines)):
        line = lines[index]
        if line.startswith(">"):
            meta.append(line[1:].strip())
        elif line.strip():
            start = index
            break
        else:
            start = index + 1

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)
    set_east_asian(normal, body_font)

    for name, size, color, before, after in (
        ("Heading 1", 16, "1F3864", 16, 8),
        ("Heading 2", 13, "2E5496", 12, 6),
        ("Heading 3", 11.5, "333333", 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        set_east_asian(style, head_font)

    # ---- cover page ----
    for _ in range(4):
        doc.add_paragraph()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.font.size = Pt(26)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    set_east_asian(run, head_font)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(subtitle)
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(0x2E, 0x54, 0x96)
    set_east_asian(run, head_font)

    doc.add_paragraph()
    doc.add_paragraph()
    for item in meta:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(paragraph, item, size=10, color="595959",
                 code_font=code_font, east_asian=body_font)

    body = parse_markdown(lines, start)

    # ---- table of contents ----
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = True
    run = paragraph.add_run("目　　录")
    run.font.size = Pt(18)
    run.bold = True
    set_east_asian(run, head_font)
    doc.add_paragraph()

    for item in body:
        if item[0] == "h" and item[1] in (2, 3):
            level, text = item[1], item[2]
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.left_indent = Cm(0 if level == 2 else 0.8)
            run = paragraph.add_run(text)
            run.font.size = Pt(10.5 if level == 2 else 10)
            run.bold = level == 2
            if level == 3:
                run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            set_east_asian(run, body_font)

    def render_table(rows):
        ncol = max(len(row) for row in rows)
        table = doc.add_table(rows=len(rows), cols=ncol)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        tblpr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:" + edge)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "BFBFBF")
            borders.append(el)
        tblpr.insert_element_before(borders, "w:shd", "w:tblLayout",
                                    "w:tblCellMar", "w:tblLook")
        tbl_w = tblpr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tblpr.insert_element_before(tbl_w, "w:jc", "w:tblInd", "w:tblBorders",
                                        "w:shd", "w:tblLayout")
        tbl_w.set(qn("w:type"), "pct")
        tbl_w.set(qn("w:w"), "5000")

        total_cm = 16.2
        weights = []
        for col in range(ncol):
            values = [text_width(row[col]) for row in rows
                      if col < len(row) and row[col].strip()]
            average = sum(values) / len(values) if values else 8
            weights.append(max(6.0, min(average, 60.0)))
        width_sum = sum(weights)
        widths = [max(1.5, total_cm * weight / width_sum) for weight in weights]
        scale = total_cm / sum(widths)
        widths = [width * scale for width in widths]
        for index, width in enumerate(widths):
            table.columns[index].width = Cm(width)

        for row_index, row in enumerate(rows):
            trpr = table.rows[row_index]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for col_index in range(ncol):
                cell = table.cell(row_index, col_index)
                cell.width = Cm(widths[col_index])
                text = row[col_index] if col_index < len(row) else ""
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.15
                add_runs(paragraph, text, size=8.8, code_font=code_font,
                         east_asian=body_font)
                if row_index == 0:
                    shade_cell(cell, "E7EEF7")
                    for run in paragraph.runs:
                        run.bold = True
        doc.add_paragraph()

    def render_code(code_lines):
        paragraph = doc.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(4)
        fmt.space_after = Pt(8)
        fmt.line_spacing = 1.0
        fmt.left_indent = Cm(0.2)
        fmt.right_indent = Cm(0.2)
        shade_paragraph(paragraph, "F4F6F8")
        for index, line in enumerate(code_lines):
            run = paragraph.add_run(line)
            run.font.name = code_font
            run.font.size = Pt(8.3)
            set_east_asian(run, body_font)
            if index < len(code_lines) - 1:
                run.add_break()

    def render_quote(quote_lines):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(8)
        add_left_border(paragraph)
        for index, line in enumerate(quote_lines):
            run = paragraph.add_run(line)
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            set_east_asian(run, body_font)
            if index < len(quote_lines) - 1:
                run.add_break()

    first_heading = True
    for item in body:
        kind = item[0]
        if kind == "h":
            level, text = item[1], item[2]
            heading = doc.add_heading(level=max(1, level - 1))
            if first_heading:
                heading.paragraph_format.page_break_before = True
                first_heading = False
            heading.text = ""
            add_runs(heading, text, code_font=code_font, east_asian=head_font)
        elif kind == "p":
            add_runs(doc.add_paragraph(), item[1], code_font=code_font,
                     east_asian=body_font)
        elif kind == "ul":
            for entry in item[1]:
                paragraph = doc.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(2)
                add_runs(paragraph, entry, code_font=code_font, east_asian=body_font)
        elif kind == "ol":
            for entry in item[1]:
                paragraph = doc.add_paragraph(style="List Number")
                paragraph.paragraph_format.space_after = Pt(2)
                add_runs(paragraph, entry, code_font=code_font, east_asian=body_font)
        elif kind == "quote":
            render_quote(item[1])
        elif kind == "code":
            render_code(item[1])
        elif kind == "table":
            render_table(item[1])

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    paragraph.add_run(" 页")
    for run in paragraph.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.save(output_path)
    return len(body), len(doc.tables)


def main():
    parser = argparse.ArgumentParser(description="Convert Markdown to DOCX")
    parser.add_argument("input", help="source markdown file")
    parser.add_argument("output", help="target docx file")
    parser.add_argument("--body-font", default="宋体", help="east asian body font")
    parser.add_argument("--head-font", default="黑体", help="east asian heading font")
    parser.add_argument("--code-font", default="Consolas", help="monospace font")
    args = parser.parse_args()
    blocks, tables = build(args.input, args.output, args.body_font,
                           args.head_font, args.code_font)
    print("saved %s (blocks=%d, tables=%d)" % (args.output, blocks, tables))


if __name__ == "__main__":
    main()
